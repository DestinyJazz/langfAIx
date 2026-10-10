"""Spreadsheets: turn Excel, CSV or text copied from Excel into glossary rows, and export them back.

Import has two halves. This module only reads the table and works out what each column
is; whether a row may enter the glossary is decided in store.py.
"""

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook

from app.glossary import Entry, tidy
from app.glossary_profile import DEFAULT_PROFILE, Profile

FIELDS = ("id", "zh", "en", "category", "applies_to", "variants", "note")
# Column headers on export. Import recognises all of them, so an exported file re-imports as is.
EXPORT_HEADERS = {
    "id": "id",
    "zh": "Chinese",
    "en": "English",
    "category": "Category",
    "applies_to": "Applies to",
    "variants": "Wrong forms",
    "note": "Note",
}

# Keywords for recognising headers, Chinese and English: a header containing one of them
# is taken as that field. Order matters: "wrong forms" before "English" because of headers
# like "Wrong English forms", and Chinese before English because of "Chinese term"
# ("term" is also a keyword for the English column).
HEADER_KEYWORDS = (
    ("variants", ("错误", "禁用", "不要用", "wrong", "forbidden", "variant", "avoid")),
    ("applies_to", ("版本", "适用", "version", "applies")),
    ("category", ("分类", "类别", "类型", "category", "type")),
    ("note", ("说明", "备注", "描述", "note", "remark", "comment", "description")),
    ("zh", ("中文", "原文", "源词", "chinese", "source", "zh", "cn")),
    ("en", ("英文", "译文", "标准", "english", "target", "translation", "term", "standard", "en")),
)

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000

# Several wrong forms in one cell are separated by semicolons, pipes or line breaks.
_VARIANT_SPLIT = re.compile(r"[;；|\n]+")


class SheetError(Exception):
    """The file cannot be read. The message is shown to the user."""


@dataclass
class Sheet:
    columns: list[str]  # column names: the header, or Column A, Column B when there is none
    mapping: dict[str, int]  # field -> column index (0-based)
    has_header: bool
    table: list[tuple[int, list[str]]] = field(default_factory=list)  # (row number in the sheet, cells)

    def rows(self) -> list[dict]:
        """Turn each row into raw term data using the current column mapping."""
        result = []
        for line, cells in self.table:
            row: dict = {"line": line}
            for name, index in self.mapping.items():
                value = cells[index] if 0 <= index < len(cells) else ""
                if name == "variants":
                    row[name] = [tidy(piece) for piece in _VARIANT_SPLIT.split(value) if tidy(piece)]
                else:
                    row[name] = value
            result.append(row)
        return result


def _cell(value) -> str:
    """Text of one cell."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # integers in Excel sometimes read back as 8.0
    return str(value).strip()


def _decode(raw: bytes) -> str:
    """Decode a CSV file.

    On Chinese Windows, Excel often saves CSV in a GB encoding rather than UTF-8, so the
    likely encodings are tried in turn.
    """
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")  # last resort: at least the ASCII part is right


def _read_xlsx(raw: bytes) -> list[tuple[int, list[str]]]:
    try:
        # read_only: stream rows instead of loading the file. data_only: computed values, not formulas.
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many exception types for broken or non-xlsx files
        raise SheetError("This file could not be read as an Excel workbook (.xlsx).") from exc
    try:
        sheet = workbook.worksheets[0]
        table = []
        for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            texts = [_cell(value) for value in cells]
            if any(texts):
                table.append((number, texts))
            if len(table) > MAX_ROWS + 1:
                break
        return table
    finally:
        workbook.close()


def _read_text(text: str, single_column_fallback: bool) -> list[tuple[int, list[str]]]:
    """Read CSV, or text copied from Excel (tab separated)."""
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []
    first = lines[0]
    if "\t" in first:
        delimiter = "\t"
    elif single_column_fallback and "," not in first and ";" not in first:
        # The pasted text is a single column of words, one per line.
        return [(number, [line.strip()]) for number, line in enumerate(text.splitlines(), start=1) if line.strip()][: MAX_ROWS + 1]
    else:
        # Excel uses a semicolon as the CSV separator in some locales.
        delimiter = ";" if first.count(";") > first.count(",") else ","
    table = []
    for number, cells in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        texts = [cell.strip() for cell in cells]
        if any(texts):
            table.append((number, texts))
        if len(table) > MAX_ROWS + 1:
            break
    return table


def _guess_field(header: str, source_name: str = "") -> str | None:
    """Which field this header is, or None.

    source_name is the configured source language ("Russian"): a header containing it is the
    source column, as "Chinese" is for the default.
    """
    name = tidy(header).casefold()
    if not name:
        return None
    if name in ("id", "编号", "序号"):
        return "id"
    for target, keywords in HEADER_KEYWORDS:
        if target == "zh" and source_name:
            keywords = keywords + (source_name.casefold(),)
        for keyword in keywords:
            # Very short ASCII keywords (zh, cn, en) must be the whole header, otherwise the "en"
            # in "content" or "Recommended" would be taken for the English column.
            if len(keyword) <= 2 and keyword.isascii():
                if name == keyword:
                    return target
            elif keyword in name:
                return target
    return None


def _detect(table: list[tuple[int, list[str]]], profile: Profile) -> tuple[dict[str, int], bool]:
    """Work out what each column is and whether the first row is a header."""
    first = table[0][1]
    mapping: dict[str, int] = {}
    for index, header in enumerate(first):
        target = _guess_field(header, profile.source_name)
        if target and target not in mapping:
            mapping[target] = index
    if "en" in mapping or "zh" in mapping:
        return mapping, True

    # No header. Guess from the content: a column that is mostly characters of the source
    # script (Han by default) is the source language, the first other column is English.
    # Without a source script there is no such column.
    source_char = re.compile(f"[{profile.chars}]") if profile.has_source else None
    width = max(len(cells) for _, cells in table)
    sample = table[:50]
    mapping = {}
    for index in range(width):
        values = [cells[index] for _, cells in sample if index < len(cells) and cells[index]]
        if not values:
            continue
        chinese = bool(source_char) and sum(1 for value in values if source_char.search(value)) > len(values) / 2
        if chinese and "zh" not in mapping:
            mapping["zh"] = index
        elif not chinese and "en" not in mapping:
            mapping["en"] = index
    return mapping, False


def read_sheet(
    filename: str = "",
    raw: bytes | None = None,
    text: str | None = None,
    mapping: dict[str, int] | None = None,
    has_header: bool | None = None,
    profile: Profile | None = None,
) -> Sheet:
    """Read an uploaded file (raw) or pasted text (text).

    mapping and has_header are normally detected. When detection is wrong the user corrects
    it on the preview page, which then repeats the request with the corrected values.
    profile says what the source language is; the default is Chinese.
    """
    profile = profile or DEFAULT_PROFILE
    name = (filename or "").lower()
    if raw is not None:
        if len(raw) > MAX_BYTES:
            raise SheetError(f"The file is too large ({len(raw) // 1024} KB, the limit is {MAX_BYTES // 1024} KB).")
        if name.endswith(".xls"):
            raise SheetError("Old .xls files are not supported. Open it in Excel and save it as .xlsx or .csv.")
        # An xlsx file is a zip archive, so it always starts with PK.
        if name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
            table = _read_xlsx(raw)
        else:
            table = _read_text(_decode(raw), single_column_fallback=False)
    else:
        if len(text or "") > MAX_BYTES:
            raise SheetError("The pasted text is too long.")
        table = _read_text(text or "", single_column_fallback=True)

    if not table:
        raise SheetError("There is nothing to import: the file or the pasted text is empty.")
    detected_mapping, detected_header = _detect(table, profile)
    if has_header is None:
        has_header = detected_header
    if mapping is None:
        mapping = detected_mapping
    else:
        # Supplied by the user: keep known fields only, with a non-negative integer column index.
        mapping = {name: int(index) for name, index in mapping.items() if name in FIELDS and isinstance(index, int) and index >= 0}
    width = max(len(cells) for _, cells in table)
    headers = table[0][1] if has_header else []
    columns = [(headers[index] if index < len(headers) and headers[index] else f"Column {_letter(index)}") for index in range(width)]
    body = table[1:] if has_header else table
    if len(body) > MAX_ROWS:
        raise SheetError(f"Too many rows (the limit is {MAX_ROWS}). Split the file and import it in parts.")
    return Sheet(columns=columns, mapping=mapping, has_header=has_header, table=body)


def _letter(index: int) -> str:
    """Excel column name for an index: 0 is A, 25 is Z, 26 is AA."""
    name = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(ord("A") + remainder) + name
    return name


# ---------- Export ----------


def _export_rows(entries: list[Entry], profile: Profile) -> list[list]:
    # The source column is headed with the source language's name, which import recognises again.
    headers = {**EXPORT_HEADERS, "zh": profile.source_name}
    rows: list[list] = [[headers[name] for name in FIELDS]]
    for entry in entries:
        rows.append([entry.id, entry.zh, entry.en, entry.category, entry.applies_to, "; ".join(entry.variants), entry.note])
    return rows


def to_csv(entries: list[Entry], profile: Profile | None = None) -> bytes:
    """Export as CSV, with a BOM so Excel reads it as UTF-8."""
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\r\n").writerows(_export_rows(entries, profile or DEFAULT_PROFILE))
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(entries: list[Entry], profile: Profile | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Glossary"
    for row in _export_rows(entries, profile or DEFAULT_PROFILE):
        sheet.append(row)
    sheet.freeze_panes = "A2"  # keep the header row in place when scrolling
    for letter, width in zip("ABCDEFG", (8, 28, 34, 16, 20, 34, 50)):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
