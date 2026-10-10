"""Reading and writing spreadsheets: Excel, CSV in several encodings, text copied from Excel, and export."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.glossary import Entry
from app.sheets import SheetError, read_sheet, to_csv, to_xlsx


def xlsx(rows: list[list]) -> bytes:
    """Make a real .xlsx file."""
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


CHINESE_HEADERS = [
    ["中文", "英文", "分类", "适用版本", "错误写法", "备注"],
    ["网页管理界面", "WebUI", "Feature", "", "Web UI; web-ui", "主界面"],
    ["归档中心", "Archive Center", "Feature", "3.1.60 及以后", "", ""],
]


def test_an_excel_file_with_chinese_headers():
    sheet = read_sheet("glossary.xlsx", xlsx(CHINESE_HEADERS))

    assert sheet.has_header
    assert sheet.columns == ["中文", "英文", "分类", "适用版本", "错误写法", "备注"]
    assert sheet.mapping == {"zh": 0, "en": 1, "category": 2, "applies_to": 3, "variants": 4, "note": 5}
    first, second = sheet.rows()
    assert first == {"line": 2, "zh": "网页管理界面", "en": "WebUI", "category": "Feature", "applies_to": "", "variants": ["Web UI", "web-ui"], "note": "主界面"}
    assert second["applies_to"] == "3.1.60 及以后" and second["variants"] == []


def test_english_headers_in_any_order():
    table = [["English Term", "Chinese", "Version", "Wrong forms", "Notes", "Type", "id"], ["WebUI", "网页管理界面", "8.0", "Web UI", "x", "Feature", 7]]
    sheet = read_sheet("g.xlsx", xlsx(table))
    assert sheet.mapping == {"en": 0, "zh": 1, "applies_to": 2, "variants": 3, "note": 4, "category": 5, "id": 6}
    assert sheet.rows()[0]["id"] == "7"


def test_headers_that_mention_two_things_are_read_sensibly():
    # "Chinese term" is the Chinese column (despite "term"); "Wrong English forms" is wrong forms (despite "English").
    sheet = read_sheet("g.xlsx", xlsx([["Chinese term", "English term", "Wrong English forms"], ["策略", "Policies", "Policys"]]))
    assert sheet.mapping == {"zh": 0, "en": 1, "variants": 2}


def test_short_keywords_do_not_match_inside_other_words():
    # "Content" and "Recommended" contain "en" but are not the English column.
    sheet = read_sheet("g.xlsx", xlsx([["Content", "English"], ["x", "WebUI"]]))
    assert sheet.mapping == {"en": 1}


def test_numbers_and_empty_rows_in_excel():
    sheet = read_sheet("g.xlsx", xlsx([["English", "Version"], ["AF", 8.0], [None, None], ["BF", 8.5], ["CF", 13]]))
    assert [(line, row["en"], row["applies_to"]) for line, row in ((r["line"], r) for r in sheet.rows())] == [
        (2, "AF", "8"),  # an integer may come out of Excel as 8.0; written back as 8
        (4, "BF", "8.5"),  # the empty row is skipped, but row numbers are still the sheet's
        (5, "CF", "13"),
    ]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gb18030", "utf-16"])
def test_csv_in_the_encodings_excel_produces(encoding):
    # On Chinese Windows, Excel often saves CSV in a GB encoding. Reading it as UTF-8 garbles the Chinese, so encodings are tried in turn.
    text = "中文,English,备注\r\n网页管理界面,WebUI,主界面\r\n策略,Policies,\r\n"
    sheet = read_sheet("glossary.csv", text.encode(encoding))
    assert sheet.mapping == {"zh": 0, "en": 1, "note": 2}
    assert [(row["zh"], row["en"]) for row in sheet.rows()] == [("网页管理界面", "WebUI"), ("策略", "Policies")]


def test_csv_with_semicolons_and_quoted_cells():
    text = 'English;Chinese;Note\r\n"Zyphor Gateway";网关;"has; a semicolon"\r\n'
    (row,) = read_sheet("g.csv", text.encode("utf-8")).rows()
    assert (row["en"], row["zh"], row["note"]) == ("Zyphor Gateway", "网关", "has; a semicolon")


def test_text_pasted_from_excel_is_tab_separated():
    text = "中文\tEnglish\n网页管理界面\tWebUI\n策略\tPolicies\n"
    sheet = read_sheet(text=text)
    assert sheet.has_header and sheet.mapping == {"zh": 0, "en": 1}
    assert [row["en"] for row in sheet.rows()] == ["WebUI", "Policies"]


def test_a_pasted_list_of_words_is_one_column_of_english():
    sheet = read_sheet(text="WebUI\nZyphor Gateway\n\nZX3.1.60\n")
    assert not sheet.has_header
    assert sheet.columns == ["Column A"]
    assert [(row["line"], row["en"]) for row in sheet.rows()] == [(1, "WebUI"), (2, "Zyphor Gateway"), (4, "ZX3.1.60")]


def test_without_headers_the_chinese_column_is_guessed_from_its_content():
    sheet = read_sheet("g.xlsx", xlsx([["网页管理界面", "WebUI"], ["策略", "Policies"], ["REST API 网关", "REST API gateway"]]))
    assert not sheet.has_header
    assert sheet.mapping == {"zh": 0, "en": 1}
    assert len(sheet.rows()) == 3  # the first row is data too, not dropped as a header


def test_the_user_can_correct_what_was_guessed():
    # Unrecognised headers: the program assumes there is no header row. The user then states on the page: the first row is a header, column A is English, column C is Chinese.
    raw = xlsx([["Name", "Owner", "Original"], ["WebUI", "amy", "网页管理界面"]])
    guessed = read_sheet("g.xlsx", raw)
    assert not guessed.has_header

    fixed = read_sheet("g.xlsx", raw, mapping={"en": 0, "zh": 2, "bogus": 1}, has_header=True)
    assert fixed.mapping == {"en": 0, "zh": 2}  # unknown field names are dropped
    assert fixed.columns == ["Name", "Owner", "Original"]
    assert fixed.rows() == [{"line": 2, "en": "WebUI", "zh": "网页管理界面"}]


def test_wrong_forms_in_one_cell_are_split():
    sheet = read_sheet("g.xlsx", xlsx([["English", "Wrong forms"], ["WebUI", "Web UI；web-ui | Web-UI\nwebui;"]]))
    assert sheet.rows()[0]["variants"] == ["Web UI", "web-ui", "Web-UI", "webui"]


def test_files_that_cannot_be_read_say_why():
    with pytest.raises(SheetError, match="save it as .xlsx or .csv"):
        read_sheet("old.xls", b"\xd0\xcf\x11\xe0 old excel")
    with pytest.raises(SheetError, match="could not be read as an Excel workbook"):
        read_sheet("broken.xlsx", b"PK this is not really a workbook")
    with pytest.raises(SheetError, match="nothing to import"):
        read_sheet("empty.csv", b"\r\n\r\n")
    with pytest.raises(SheetError, match="nothing to import"):
        read_sheet(text="   ")
    with pytest.raises(SheetError, match="too large"):
        read_sheet("big.csv", b"x" * (5 * 1024 * 1024 + 1))


# ---------- Export ----------

ENTRIES = [
    Entry(id=3, en="WebUI", zh="网页管理界面", category="Feature", variants=("Web UI", "web-ui"), note="主界面"),
    Entry(id=4, en="Archive Center", zh="归档中心", applies_to="3.1.60 及以后"),
]


def test_exported_csv_opens_in_excel_and_can_be_imported_again():
    raw = to_csv(ENTRIES)
    assert raw.startswith(b"\xef\xbb\xbf")  # the BOM tells Excel this is UTF-8

    sheet = read_sheet("glossary.csv", raw)
    assert sheet.mapping == {"id": 0, "zh": 1, "en": 2, "category": 3, "applies_to": 4, "variants": 5, "note": 6}
    first, second = sheet.rows()
    assert first == {"line": 2, "id": "3", "zh": "网页管理界面", "en": "WebUI", "category": "Feature", "applies_to": "", "variants": ["Web UI", "web-ui"], "note": "主界面"}
    assert (second["id"], second["applies_to"]) == ("4", "3.1.60 及以后")


def test_exported_excel_can_be_imported_again():
    raw = to_xlsx(ENTRIES)
    workbook = load_workbook(io.BytesIO(raw))
    assert [cell.value for cell in workbook.active[1]] == ["id", "Chinese", "English", "Category", "Applies to", "Wrong forms", "Note"]

    rows = read_sheet("glossary.xlsx", raw).rows()
    assert [(row["id"], row["zh"], row["en"], row["variants"]) for row in rows] == [
        ("3", "网页管理界面", "WebUI", ["Web UI", "web-ui"]),
        ("4", "归档中心", "Archive Center", []),
    ]

# ---------- Another source language ----------


def test_the_source_language_name_is_recognised_as_a_header_and_used_when_exporting():
    from app.glossary_profile import make_profile

    profile = make_profile("Russian", "cyrillic")
    sheet = read_sheet("g.xlsx", xlsx([["Russian", "English"], ["Политики", "Policies"]]), profile=profile)
    assert sheet.mapping == {"zh": 0, "en": 1}
    # The same header is not the source column in a default (Chinese) glossary.
    assert read_sheet("g.xlsx", xlsx([["Russian", "English"], ["Политики", "Policies"]])).mapping == {"en": 1}

    exported = to_csv([Entry(id=1, en="Policies", zh="Политики")], profile).decode("utf-8-sig")
    assert exported.splitlines()[0].split(",")[:3] == ["id", "Russian", "English"]
    # It reads back as it was written.
    assert read_sheet("g.csv", to_csv([Entry(id=1, en="Policies", zh="Политики")], profile), profile=profile).mapping["zh"] == 1


def test_without_headers_the_source_column_is_found_by_the_configured_script():
    from app.glossary_profile import make_profile

    table = [["Политики", "Policies"], ["Файлы", "Files"]]
    sheet = read_sheet("g.xlsx", xlsx(table), profile=make_profile("Russian", "cyrillic"))
    assert sheet.mapping == {"zh": 0, "en": 1} and not sheet.has_header
    # With no source script, no column is the source column.
    assert "zh" not in read_sheet("g.xlsx", xlsx(table), profile=make_profile(script="none")).mapping