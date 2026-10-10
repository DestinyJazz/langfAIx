"""The glossary endpoints, from the HTTP request in to the HTTP response out.

LiteLLM is a fake (it only identifies keys); the database is an in-memory SQLite.
All terms in the examples are made up.
"""

import base64
import io

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from app.identity import hash_key
from tests.conftest import ADMIN_KEY, LLM_URL, FakeKeys, lt_spelling

# Three people: root is the admin named in the environment, amy and ben are ordinary users.
KEYS = {"sk-root": "root", "sk-amy": "amy", "sk-ben": "ben"}
ROOT = {"Authorization": "Bearer sk-root"}
AMY = {"Authorization": "Bearer sk-amy"}
BEN = {"Authorization": "Bearer sk-ben"}

WEBUI = {"en": "WebUI", "zh": "网页管理界面", "category": "Feature", "variants": ["Web UI"], "note": "main page"}


@pytest.fixture
def client(make_client):
    return make_client(FakeKeys(KEYS), glossary_admins="root", litellm_admin_key=ADMIN_KEY)


def make_editor(client, headers, key: str) -> None:
    """Let someone sign in once, then have the admin make them an editor: the real granting flow."""
    assert client.get("/glossary/me", headers=headers).status_code == 200
    assert client.put(f"/glossary/people/{hash_key(key)}", json={"role": "editor"}, headers=ROOT).status_code == 200


# ---------- Identity ----------


def test_every_glossary_endpoint_needs_a_valid_key(client):
    for method, path in [("GET", "/glossary/me"), ("GET", "/glossary/terms"), ("GET", "/glossary/lookup?q=x"), ("POST", "/glossary/terms"), ("GET", "/glossary/export")]:
        missing = client.request(method, path, json={})
        assert missing.status_code == 401, path
        assert "missing API key" in missing.json()["detail"]
        wrong = client.request(method, path, json={}, headers={"Authorization": "Bearer sk-nobody"})
        assert wrong.status_code == 401, path
        assert "rejected" in wrong.json()["detail"]


def test_me_says_who_you_are_and_what_you_may_do(client):
    amy = client.get("/glossary/me", headers=AMY).json()
    assert amy == {
        "alias": "amy",
        "id": hash_key("sk-amy")[:12],
        "role": "viewer",
        "can_edit": False,
        "is_admin": False,
        # The next four belong to password sign-in; see test_login_api.py.
        "login": "",
        "has_password": False,
        "signed_in_with": "key",
        "password_sign_in": True,
    }

    root = client.get("/glossary/me", headers=ROOT).json()
    assert (root["role"], root["can_edit"], root["is_admin"]) == ("admin", True, True)


def test_user_keys_are_looked_up_by_the_server_with_its_own_key(make_client):
    # The usual deployment: a user's key may only call models and cannot look itself up. So
    # the backend asks with its own admin key, by hash. The real key must never be in a URL
    # (URLs end up in access logs).
    litellm = FakeKeys(KEYS)
    client = make_client(litellm, glossary_admins="root", litellm_admin_key=ADMIN_KEY)

    assert client.get("/glossary/me", headers=AMY).json()["alias"] == "amy"

    (request,) = litellm.requests
    assert request.headers["authorization"] == f"Bearer {ADMIN_KEY}"
    assert request.url.params["key"] == hash_key("sk-amy")
    assert "sk-amy" not in str(request.url)


def test_without_the_server_key_a_model_only_key_gets_a_clear_explanation(make_client):
    # This once produced a bare "key rejected", which made people think the key was mistyped.
    client = make_client(FakeKeys(KEYS), glossary_admins="root")
    response = client.get("/glossary/me", headers=AMY)
    assert response.status_code == 503
    assert "LITELLM_ADMIN_KEY" in response.json()["detail"]
    assert "not a problem with your key" in response.json()["detail"]


def test_a_wrong_server_key_is_not_blamed_on_the_user(make_client):
    client = make_client(FakeKeys(KEYS), glossary_admins="root", litellm_admin_key="sk-this-is-not-the-master-key")
    response = client.get("/glossary/me", headers=AMY)
    assert response.status_code == 502
    assert "LITELLM_ADMIN_KEY" in response.json()["detail"]


def test_default_type_keys_can_still_identify_themselves(make_client):
    # No admin key configured and default-type keys: a key looks itself up, and that works.
    litellm = FakeKeys(KEYS, restricted=False)
    client = make_client(litellm, glossary_admins="root")
    assert client.get("/glossary/me", headers=ROOT).json()["role"] == "admin"
    assert str(litellm.requests[0].url) == f"{LLM_URL}/key/info"


def test_litellm_being_down_is_reported_as_such(make_client):
    def litellm_down(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("refused")

    response = make_client(litellm_down).get("/glossary/me", headers=AMY)
    assert response.status_code == 502
    assert "LiteLLM" in response.json()["detail"]


# ---------- Roles ----------


def test_a_viewer_can_read_but_not_write(client):
    client.post("/glossary/terms", json=WEBUI, headers=ROOT)

    assert client.get("/glossary/terms", headers=AMY).status_code == 200
    assert client.get("/glossary/lookup?q=webui", headers=AMY).status_code == 200
    assert client.get("/glossary/export", headers=AMY).status_code == 200
    assert client.get("/glossary/changes", headers=AMY).status_code == 200

    for response in (
        client.post("/glossary/terms", json={"en": "Zyphor"}, headers=AMY),
        client.put("/glossary/terms/1", json={"en": "Zyphor"}, headers=AMY),
        client.delete("/glossary/terms/1", headers=AMY),
        client.post("/glossary/import", json={"rows": [{"en": "Zyphor"}]}, headers=AMY),
        client.post("/glossary/import/preview", json={"text": "Zyphor"}, headers=AMY),
    ):
        assert response.status_code == 403
        assert "needs the editor role" in response.json()["detail"]
    # Nothing was changed.
    assert [term["en"] for term in client.get("/glossary/terms", headers=AMY).json()["terms"]] == ["WebUI"]


def test_an_admin_makes_someone_an_editor(client):
    make_editor(client, AMY, "sk-amy")

    assert client.get("/glossary/me", headers=AMY).json()["role"] == "editor"
    created = client.post("/glossary/terms", json=WEBUI, headers=AMY)
    assert created.status_code == 201
    assert created.json()["created_by"] == "amy"
    # ben was not granted anything and still cannot write.
    assert client.post("/glossary/terms", json={"en": "Zyphor"}, headers=BEN).status_code == 403


def test_only_admins_manage_people(client):
    make_editor(client, AMY, "sk-amy")
    client.get("/glossary/me", headers=BEN)

    # An editor cannot see the people list or change anyone's role, their own included.
    assert client.get("/glossary/people", headers=AMY).status_code == 403
    assert client.put(f"/glossary/people/{hash_key('sk-amy')}", json={"role": "admin"}, headers=AMY).status_code == 403

    people = {person["alias"]: person for person in client.get("/glossary/people", headers=ROOT).json()["people"]}
    assert {name: person["role"] for name, person in people.items()} == {"amy": "editor", "ben": "viewer"}
    assert all("key_hash" in person and not person["fixed"] for person in people.values())

    assert client.put(f"/glossary/people/{hash_key('sk-ben')}", json={"role": "owner"}, headers=ROOT).status_code == 422
    assert client.put("/glossary/people/unknownhash", json={"role": "editor"}, headers=ROOT).status_code == 404


def test_admins_from_the_setting_cannot_be_demoted_on_the_page(client):
    client.get("/glossary/me", headers=ROOT)
    client.put(f"/glossary/people/{hash_key('sk-root')}", json={"role": "viewer"}, headers=ROOT)

    (person,) = client.get("/glossary/people", headers=ROOT).json()["people"]
    assert (person["alias"], person["role"], person["fixed"]) == ("root", "admin", True)
    assert client.get("/glossary/me", headers=ROOT).json()["is_admin"] is True


# ---------- Add, change, delete ----------


def test_add_list_update_delete(client):
    created = client.post("/glossary/terms", json=WEBUI, headers=ROOT)
    assert created.status_code == 201
    term = created.json()
    assert term["id"] > 0
    assert {key: term[key] for key in WEBUI} == WEBUI
    assert (term["created_by"], term["updated_by"]) == ("root", "root")

    listing = client.get("/glossary/terms", headers=ROOT).json()
    assert [item["en"] for item in listing["terms"]] == ["WebUI"]
    assert "Version" in listing["categories"]

    changed = client.put(f"/glossary/terms/{term['id']}", json={**WEBUI, "applies_to": "8.0", "category": "Portal"}, headers=ROOT)
    assert changed.status_code == 200
    assert changed.json()["applies_to"] == "8.0"
    assert "Portal" in client.get("/glossary/terms", headers=ROOT).json()["categories"]  # a new category appears once it is used

    assert client.delete(f"/glossary/terms/{term['id']}", headers=ROOT).status_code == 204
    assert client.get("/glossary/terms", headers=ROOT).json()["terms"] == []


def test_input_is_tidied_and_problems_are_explained(client):
    messy = client.post("/glossary/terms", json={"en": "  Zyphor   Gateway ", "zh": " 网关 ", "variants": ["Zyfor Gateway", "", "Zyfor Gateway"]}, headers=ROOT).json()
    assert (messy["en"], messy["zh"], messy["variants"]) == ("Zyphor Gateway", "网关", ["Zyfor Gateway"])

    empty = client.post("/glossary/terms", json={"zh": "只有中文"}, headers=ROOT)
    assert empty.status_code == 422
    assert empty.json()["detail"] == "The English term is required."

    duplicate = client.post("/glossary/terms", json={"en": "zyphor gateway", "zh": "网关"}, headers=ROOT)
    assert duplicate.status_code == 422
    assert "already exists" in duplicate.json()["detail"]


def test_missing_rows_give_404(client):
    assert client.put("/glossary/terms/99", json={"en": "X"}, headers=ROOT).status_code == 404
    assert client.delete("/glossary/terms/99", headers=ROOT).status_code == 404


def test_changes_are_logged(client):
    make_editor(client, AMY, "sk-amy")
    term = client.post("/glossary/terms", json=WEBUI, headers=AMY).json()
    client.delete(f"/glossary/terms/{term['id']}", headers=ROOT)

    log = client.get("/glossary/changes", headers=BEN).json()["changes"]
    assert [(item["action"], item["by"]) for item in log] == [("delete", "root"), ("add", "amy"), ("role", "root")]
    assert log[0]["details"]["before"]["en"] == "WebUI"


# ---------- Changes take effect at once ----------


def test_a_new_term_takes_effect_on_the_next_check(make_client):
    text = "Open the WebUI and the Web UI, then click 新增 to continue."
    lt_reply = {"matches": [lt_spelling(text, "WebUI", ("Web UI",))]}
    client = make_client(FakeKeys(KEYS, lt_reply), glossary_admins="root", litellm_admin_key=ADMIN_KEY)

    before = client.post("/check", json={"text": text}).json()["matches"]
    assert [match["text"] for match in before] == ["WebUI"]  # the glossary is still empty: LanguageTool reports it

    client.post("/glossary/terms", json=WEBUI, headers=ROOT)
    client.post("/glossary/terms", json={"en": "Add", "zh": "新增"}, headers=ROOT)

    # No restart: the next check uses the new glossary.
    after = client.post("/check", json={"text": text}).json()["matches"]
    assert [(match["text"], match["replacements"], match["issue_type"]) for match in after] == [
        ("Web UI", ["WebUI"], "terminology"),  # a registered wrong form
        ("新增", ["Add"], "terminology"),  # Chinese left in English text
    ]
    assert client.get("/healthz").json()["glossary_terms"] == 2


def test_the_glossary_can_be_switched_off_for_one_request(make_client):
    # The user turned the glossary off in the extension settings: this check behaves as if there were none.
    text = "Open the WebUI and the Web UI, then click 新增 to continue."
    lt_reply = {"matches": [lt_spelling(text, "WebUI", ("Web UI",))]}
    terms = [WEBUI, {"en": "Add", "zh": "新增"}]
    client = make_client(FakeKeys(KEYS, lt_reply), terms=terms)

    def check(**extra):
        return [match["text"] for match in client.post("/check", json={"text": text, **extra}).json()["matches"]]

    # Not specified means on: terms are not misspellings; the wrong form and the stray Chinese are flagged.
    assert check() == ["Web UI", "新增"]
    assert check(glossary=True) == ["Web UI", "新增"]
    # Off: LanguageTool reports the term as a misspelling again, and the glossary's two issues are gone.
    assert check(glossary=False) == ["WebUI"]


def test_check_itself_needs_no_key(client):
    client.post("/glossary/terms", json=WEBUI, headers=ROOT)
    assert client.post("/check", json={"text": "Open the Web UI now please."}).status_code == 200


# ---------- Lookup ----------


def test_lookup_by_chinese_or_english(client):
    for term in (
        {"en": "Archiving Center", "zh": "归档中心", "applies_to": "before 3.1.60"},
        {"en": "Archive Center", "zh": "归档中心", "applies_to": "3.1.60 and later"},
        {"en": "Policies", "zh": "策略"},
    ):
        assert client.post("/glossary/terms", json=term, headers=ROOT).status_code == 201

    found = client.get("/glossary/lookup", params={"q": "归档中心"}, headers=AMY).json()["results"]
    assert [(item["en"], item["applies_to"]) for item in found] == [("Archive Center", "3.1.60 and later"), ("Archiving Center", "before 3.1.60")]
    assert [item["en"] for item in client.get("/glossary/lookup", params={"q": "polic"}, headers=AMY).json()["results"]] == ["Policies"]
    assert client.get("/glossary/lookup", params={"q": ""}, headers=AMY).json()["results"] == []


# ---------- Import and export ----------


def b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode("ascii")


def workbook(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


SHEET = [
    ["中文", "英文", "分类", "适用版本", "错误写法", "备注"],
    ["网页管理界面", "WebUI", "Feature", "", "Web UI", ""],
    ["归档中心", "Archive Center", "Feature", "3.1.60 及以后", "", "renamed"],
    ["没有英文的一行", "", "", "", "", ""],
]


def test_preview_then_import_an_excel_file(client):
    client.post("/glossary/terms", json={"en": "WebUI", "zh": "网页管理界面"}, headers=ROOT)
    body = {"filename": "glossary.xlsx", "content_base64": b64(workbook(SHEET))}

    preview = client.post("/glossary/import/preview", json=body, headers=ROOT)
    assert preview.status_code == 200
    data = preview.json()
    assert data["columns"] == SHEET[0]
    assert data["mapping"] == {"zh": 0, "en": 1, "category": 2, "applies_to": 3, "variants": 4, "note": 5}
    assert [(row["line"], row["status"]) for row in data["rows"]] == [(2, "update"), (3, "new"), (4, "error")]
    assert data["counts"] == {"new": 1, "update": 1, "same": 0, "skip": 0, "error": 1}
    assert data["rows"][2]["messages"] == ["The English term is required."]
    # A preview writes nothing.
    assert len(client.get("/glossary/terms", headers=ROOT).json()["terms"]) == 1

    done = client.post("/glossary/import", json={"rows": [row["data"] | {"line": row["line"]} for row in data["rows"]]}, headers=ROOT)
    assert done.status_code == 200
    assert done.json()["counts"] == {"new": 1, "update": 1, "same": 0, "skip": 0, "error": 1}
    assert [problem["line"] for problem in done.json()["problems"]] == [4]

    terms = {term["en"]: term for term in client.get("/glossary/terms", headers=ROOT).json()["terms"]}
    assert terms["WebUI"]["variants"] == ["Web UI"] and terms["WebUI"]["category"] == "Feature"
    assert terms["Archive Center"]["applies_to"] == "3.1.60 及以后"


def test_preview_in_add_only_mode_leaves_existing_rows_alone(client):
    client.post("/glossary/terms", json={"en": "WebUI", "zh": "网页管理界面"}, headers=ROOT)
    body = {"filename": "glossary.xlsx", "content_base64": b64(workbook(SHEET)), "update_existing": False}
    statuses = [row["status"] for row in client.post("/glossary/import/preview", json=body, headers=ROOT).json()["rows"]]
    assert statuses == ["skip", "new", "error"]


def test_preview_of_pasted_text_and_gb_encoded_csv(client):
    pasted = client.post("/glossary/import/preview", json={"text": "中文\tEnglish\n策略\tPolicies\n"}, headers=ROOT).json()
    assert [(row["data"]["zh"], row["data"]["en"], row["status"]) for row in pasted["rows"]] == [("策略", "Policies", "new")]

    csv_bytes = "中文,English\r\n策略,Policies\r\n".encode("gb18030")
    from_csv = client.post("/glossary/import/preview", json={"filename": "g.csv", "content_base64": b64(csv_bytes)}, headers=ROOT).json()
    assert from_csv["rows"][0]["data"]["zh"] == "策略"  # the Chinese was not garbled


def test_preview_asks_for_the_columns_when_it_cannot_tell(client):
    # Headers the program does not know, and both columns are Chinese: no telling which one is English.
    raw = workbook([["名称", "叫法"], ["网页管理界面", "管理页"], ["策略", "规则"]])
    data = client.post("/glossary/import/preview", json={"filename": "g.xlsx", "content_base64": b64(raw)}, headers=ROOT).json()
    assert data["needs_mapping"] is True and data["rows"] == []
    assert data["columns"] == ["Column A", "Column B"]

    # Again, after the user stated on the page: first row is a header, column A is Chinese, column B is English.
    fixed = client.post(
        "/glossary/import/preview",
        json={"filename": "g.xlsx", "content_base64": b64(raw), "mapping": {"zh": 0, "en": 1}, "has_header": True},
        headers=ROOT,
    ).json()
    assert fixed["needs_mapping"] is False
    assert fixed["columns"] == ["名称", "叫法"]
    # Here the "English" column is really Chinese too. Every row is rejected with a reason instead of Chinese being stored as English.
    assert [row["status"] for row in fixed["rows"]] == ["error", "error"]
    assert "contains Chinese characters" in fixed["rows"][0]["messages"][0]


def test_bad_uploads_are_explained(client):
    not_base64 = client.post("/glossary/import/preview", json={"filename": "g.xlsx", "content_base64": "***"}, headers=ROOT)
    assert not_base64.status_code == 400
    old_excel = client.post("/glossary/import/preview", json={"filename": "old.xls", "content_base64": b64(b"old")}, headers=ROOT)
    assert old_excel.status_code == 400 and ".xlsx or .csv" in old_excel.json()["detail"]
    nothing = client.post("/glossary/import/preview", json={"text": "  "}, headers=ROOT)
    assert nothing.status_code == 400 and "nothing to import" in nothing.json()["detail"]


def test_import_checks_every_row_again(client):
    # Rows sent back by the page are not trusted: the second duplicates the first, the third has no English.
    rows = [{"en": "WebUI", "zh": "网页管理界面"}, {"en": "webui", "zh": "网页管理界面"}, {"zh": "只有中文"}]
    done = client.post("/glossary/import", json={"rows": rows}, headers=ROOT).json()
    assert done["counts"]["new"] == 1 and done["counts"]["error"] == 1
    assert len(client.get("/glossary/terms", headers=ROOT).json()["terms"]) == 1


def test_export_then_import_round_trip(client):
    for term in (WEBUI, {"en": "Archive Center", "zh": "归档中心", "applies_to": "3.1.60 及以后"}):
        client.post("/glossary/terms", json=term, headers=ROOT)

    exported = client.get("/glossary/export", params={"format": "xlsx"}, headers=AMY)
    assert exported.status_code == 200
    assert "glossary.xlsx" in exported.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [cell.value for cell in sheet[1]] == ["id", "Chinese", "English", "Category", "Applies to", "Wrong forms", "Note"]
    assert sheet.max_row == 3

    as_csv = client.get("/glossary/export", params={"format": "csv"}, headers=AMY)
    assert as_csv.content.startswith(b"\xef\xbb\xbf")
    assert "网页管理界面" in as_csv.content.decode("utf-8-sig")
    assert client.get("/glossary/export", params={"format": "pdf"}, headers=AMY).status_code == 400

    # Re-importing an exported file as is: every row is "no change".
    preview = client.post("/glossary/import/preview", json={"filename": "glossary.xlsx", "content_base64": b64(exported.content)}, headers=ROOT).json()
    assert [row["status"] for row in preview["rows"]] == ["same", "same"]


# ---------- When the database is down ----------


def test_when_the_database_is_down_checking_still_works(make_client):
    # The URL points into a folder that does not exist, so the database cannot open.
    text = "Open the WebUI now please."
    lt_reply = {"matches": [lt_spelling(text, "WebUI")]}
    client = make_client(FakeKeys(KEYS, lt_reply), glossary_admins="root", litellm_admin_key=ADMIN_KEY, database_url="sqlite+pysqlite:///Z:/no/such/folder/glossary.db")

    assert client.get("/healthz").json() == {"status": "ok", "glossary_terms": 0, "glossary_storage": "unavailable"}
    # Rule checks still work, only without a glossary.
    assert [match["text"] for match in client.post("/check", json={"text": text}).json()["matches"]] == ["WebUI"]
    # The glossary endpoints explain the reason instead of failing obscurely.
    response = client.get("/glossary/terms", headers=ROOT)
    assert response.status_code == 503
    assert "not available right now" in response.json()["detail"]


# ---------- The admin page ----------


def test_the_admin_page_is_served(make_client):
    client = make_client(FakeKeys(KEYS), root_path="/api")
    page = client.get("/admin/")
    assert page.status_code == 200
    assert "text/html" in page.headers["content-type"]
    # Without the trailing slash it redirects to the URL with one. The redirect is relative, so it is right under /api too.
    redirect = client.get("/admin", follow_redirects=False)
    assert redirect.status_code in (302, 307)
    assert redirect.headers["location"] == "admin/"

# ---------- What the glossary is about ----------


def test_the_default_profile_is_chinese(client):
    assert client.get("/glossary/profile", headers=AMY).json() == {
        "source_language": "Chinese",
        "has_source": True,
        "categories": ["Product", "Feature", "Abbreviation", "Version", "Command", "Other"],
    }
    assert client.get("/glossary/profile").status_code == 401


def test_the_profile_follows_the_settings(make_client):
    client = make_client(
        FakeKeys(KEYS),
        glossary_admins="root",
        litellm_admin_key=ADMIN_KEY,
        glossary_source_language="Russian",
        glossary_source_script="cyrillic",
        glossary_categories="Hardware, Software",
    )
    assert client.get("/glossary/profile", headers=AMY).json() == {
        "source_language": "Russian",
        "has_source": True,
        "categories": ["Hardware", "Software"],
    }
    # The English column may not hold Russian, and the message names the language.
    refused = client.post("/glossary/terms", json={"en": "Политики"}, headers=ROOT)
    assert refused.status_code == 422 and "Russian" in refused.json()["detail"]


def test_a_mistyped_script_stops_the_start(make_client):
    with pytest.raises(ValueError, match="GLOSSARY_SOURCE_SCRIPT"):
        make_client(FakeKeys(KEYS), glossary_source_script="klingon")